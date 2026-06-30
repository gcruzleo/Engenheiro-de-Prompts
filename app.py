import streamlit as st
import os
import base64
import tempfile
from groq import Groq
from st_copy_to_clipboard import st_copy_to_clipboard
import PyPDF2
import docx

# Configuração da página - minimalista e focada em leitura
st.set_page_config(page_title="Arquiteto de Prompts Sênior", layout="centered")

# Constantes
MODEL_NAME = "llama-3.3-70b-versatile"

# Inicialização do Cliente Groq
@st.cache_resource
def get_groq_client():
    try:
        api_key = st.secrets["GROQ_API_KEY"]
    except Exception:
        api_key = os.environ.get("GROQ_API_KEY", "")
        
    if not api_key:
        st.error("Chave de API não configurada. Configure a variável GROQ_API_KEY nos secrets.")
        st.stop()
        
    return Groq(api_key=api_key)

client = get_groq_client()

# --- Processadores de Arquivos (Opção C) ---
def extract_text_from_pdf(file):
    pdf_reader = PyPDF2.PdfReader(file)
    text = ""
    for page in pdf_reader.pages:
        extracted = page.extract_text()
        if extracted:
            text += extracted + "\n"
    return text

def extract_text_from_docx(file):
    doc = docx.Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def process_file_attachment(uploaded_file, groq_client):
    """
    Processa arquivos PDF, DOCX, TXT, Áudio (Whisper) ou Imagens (Llama Vision).
    Retorna uma string formatada com o conteúdo para o contexto.
    """
    if uploaded_file is None:
        return ""
        
    file_type = uploaded_file.type
    file_name = uploaded_file.name
    
    try:
        # 1. Documentos de Texto
        if "pdf" in file_type:
            return f"\n--- CONTEÚDO DO PDF ({file_name}) ---\n" + extract_text_from_pdf(uploaded_file)
        
        elif "wordprocessingml.document" in file_type or "docx" in file_name:
            return f"\n--- CONTEÚDO DO DOCX ({file_name}) ---\n" + extract_text_from_docx(uploaded_file)
            
        elif "text/plain" in file_type:
            return f"\n--- CONTEÚDO DO TXT ({file_name}) ---\n" + uploaded_file.getvalue().decode("utf-8")
            
        # 2. Áudio (Whisper API)
        elif "audio" in file_type or file_name.lower().endswith(('.mp3', '.wav', '.m4a')):
            with tempfile.NamedTemporaryFile(delete=False, suffix=f".{file_name.split('.')[-1]}") as temp_audio:
                temp_audio.write(uploaded_file.getvalue())
                temp_audio_path = temp_audio.name
                
            try:
                with open(temp_audio_path, "rb") as f:
                    transcription = groq_client.audio.transcriptions.create(
                      file=(file_name, f.read()),
                      model="whisper-large-v3",
                      response_format="text"
                    )
                return f"\n--- TRANSCRIÇÃO DO ÁUDIO ({file_name}) ---\n" + transcription
            finally:
                os.remove(temp_audio_path)
                
        # 3. Imagens (Llama Vision API)
        elif "image" in file_type:
            base64_image = base64.b64encode(uploaded_file.getvalue()).decode('utf-8')
            chat_completion = groq_client.chat.completions.create(
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": "Faça uma descrição extremamente detalhada desta imagem. Descreva os objetos, as pessoas, as cores, o ambiente, o estilo de iluminação, a técnica visual e qualquer texto presente. Esta descrição será usada como contexto para eu criar um prompt generativo depois."},
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:{file_type};base64,{base64_image}",
                                }
                            }
                        ]
                    }
                ],
                model="llama-3.2-11b-vision-preview",
            )
            image_desc = chat_completion.choices[0].message.content
            return f"\n--- DESCRIÇÃO DA IMAGEM ({file_name}) ---\n[A IA de visão analisou a imagem anexada e relatou o seguinte:]\n" + image_desc
            
        else:
            st.warning(f"Formato ignorado: {file_type}. Apenas textos, áudios e imagens são suportados.")
            return ""
            
    except Exception as e:
        st.error(f"Erro ao processar {file_name}: {str(e)}")
        return ""

# --- Funções Dinâmicas de Prompt baseadas no Tipo de IA ---
def get_system_prompt_triagem(tipo_ia):
    return f"""Você é um analista investigativo de Engenharia de Prompts.
Sua missão é ler a ideia embrionária do usuário (e possivelmente o contexto de arquivos anexados que ele enviou) e formular de 2 a 3 perguntas cruciais e diretas para extrair o contexto final necessário para criar um prompt otimizado para: {tipo_ia}.

As perguntas devem focar em:
- O objetivo principal e contexto da solicitação.
- Detalhes técnicos, estilísticos ou visuais (dependendo de qual IA foi escolhida).
- Restrições, público-alvo ou nível de complexidade desejado.

As perguntas devem ser curtas, claras, objetivas e numeradas.
NÃO adicione saudações, introduções ou conclusões. Apenas liste as perguntas."""

def get_system_prompt_sintese(tipo_ia):
    base_prompt = "Você é um Engenheiro de Prompts Especialista. "
    
    if "IA Geral" in tipo_ia:
        base_prompt += "Sua missão é estruturar um prompt magistral focado em extrair o máximo de desempenho de LLMs gerais (Gemini, ChatGPT, Claude). O prompt final DEVE exigir alto rigor técnico, incorporar formatação em Markdown, usar jargões técnicos apropriados, exigir fontes e ter estruturação analítica densa (a não ser que o usuário tenha pedido de forma resumida)."
    elif "Vibecoding" in tipo_ia:
        base_prompt += """Sua missão é estruturar um prompt magistral focado em extrair o máximo de desempenho de agentes autônomos de programação (Antigravity, Cursor, Cline).
O FORMATO DE SAÍDA OBRIGATÓRIO do seu prompt deve ser um documento Markdown robusto com as seguintes sessões claramente delimitadas:
1. # Contexto: Explicação clara do problema.
2. # Objetivo: O que deve ser construído.
3. # Arquitetura e Estrutura de Arquivos: Definição clara de como organizar o código.
4. # Plano de Implementação Passo-a-Passo: Instruções granulares para o agente seguir.
5. # Restrições e Regras Estritas: Exigência de código limpo (SOLID), tratamento de edge cases e proibição categórica de código preguiçoso (sem placeholders).
O prompt DEVE ser escrito como uma instrução direta para o agente de IA ler."""
    elif "Imagens" in tipo_ia:
        base_prompt += """Sua missão é estruturar um prompt magistral focado em geração de imagens (Midjourney, DALL-E, Stable Diffusion). 
O FORMATO DE SAÍDA OBRIGATÓRIO deve ser EXCLUSIVAMENTE em formato JSON (.json) válido.
Estruture o JSON com chaves em inglês contendo descritores estéticos otimizados para difusão. Exemplo de estrutura requerida:
{
  "subject": "Descrição do objeto principal",
  "environment": "Ambiente e fundo",
  "lighting": "Tipos de iluminação (ex: cinematic, volumetric)",
  "camera": "Lente, ângulo, profundidade de campo",
  "style": "Estilo artístico, renderização (ex: Unreal Engine, 8k)",
  "raw_prompt": "String única juntando os valores acima separados por vírgulas, pronta para o Midjourney"
}"""
    elif "Vídeos" in tipo_ia:
        base_prompt += """Sua missão é estruturar um prompt magistral focado em geração de vídeos (Sora, Runway, Pika). 
O FORMATO DE SAÍDA OBRIGATÓRIO deve ser EXCLUSIVAMENTE em formato JSON (.json) válido.
Estruture o JSON com chaves em inglês contendo descrições exatas de movimento. Exemplo de estrutura:
{
  "subject": "Sujeito e sua ação",
  "camera_motion": "Movimento de câmera (pan, tilt, zoom, tracking)",
  "scene_dynamics": "Física e comportamento da cena ao longo do tempo",
  "lighting_and_atmosphere": "Iluminação cinematográfica, textura, clima",
  "raw_prompt": "String única juntando os valores acima pronta para o gerador de vídeos"
}"""
        
    base_prompt += "\n\nRetorne APENAS o texto do prompt final (ou o JSON puro) pronto para uso. Não escreva nenhuma introdução, conclusão ou comentários adicionais. Use blocos de código markdown apropriados (```markdown ou ```json)."
    return base_prompt

# --- Gerenciamento de Estado (State Management) ---
if 'step' not in st.session_state:
    st.session_state.step = 1
if 'ideia_inicial' not in st.session_state:
    st.session_state.ideia_inicial = ""
if 'perguntas' not in st.session_state:
    st.session_state.perguntas = ""
if 'respostas' not in st.session_state:
    st.session_state.respostas = ""
if 'prompt_final' not in st.session_state:
    st.session_state.prompt_final = ""
if 'tipo_ia' not in st.session_state:
    st.session_state.tipo_ia = "IA Geral (Gemini, ChatGPT, Claude)"

def reset_app():
    st.session_state.step = 1
    st.session_state.ideia_inicial = ""
    st.session_state.perguntas = ""
    st.session_state.respostas = ""
    st.session_state.prompt_final = ""
    st.session_state.tipo_ia = "IA Geral (Gemini, ChatGPT, Claude)"

# --- Cabeçalho Principal ---
st.title("Arquiteto de Prompts Sênior")
st.markdown("Transforme solicitações básicas e arquivos em prompts magistrais otimizados para diversas ferramentas de IA.")
st.divider()

# ETAPA 1: TRIAGEM
if st.session_state.step == 1:
    st.subheader("Etapa 1: Triagem")
    st.markdown("Descreva a sua ideia embrionária e anexe arquivos (se houver) para iniciarmos o projeto.")
    
    opcoes_ia = [
        "IA Geral (Gemini, ChatGPT, Claude)", 
        "Vibecoding / Programação (Antigravity, Cursor, etc)", 
        "Geração de Imagens (Midjourney, DALL-E, etc)", 
        "Geração de Vídeos (Sora, Runway, etc)"
    ]
    
    tipo_ia = st.selectbox(
        "Para qual tipo de IA este prompt será destinado?",
        options=opcoes_ia,
        index=opcoes_ia.index(st.session_state.tipo_ia) if st.session_state.tipo_ia in opcoes_ia else 0
    )
    
    ideia = st.text_area("Ideia Inicial:", value=st.session_state.ideia_inicial, height=150, placeholder="Ex: Preciso de um script em python... / Uma foto baseada na imagem anexada...")
    
    uploaded_file = st.file_uploader("Anexar arquivo de contexto (Opcional - Textos, Áudios ou Imagens)", type=["pdf", "txt", "docx", "mp3", "wav", "m4a", "png", "jpg", "jpeg"])
    
    if st.button("Analisar Ideia", type="primary"):
        if ideia.strip() or uploaded_file is not None:
            ideia_completa = ideia
            
            # Se houver arquivo, o sistema "lê" e traduz para texto
            if uploaded_file is not None:
                with st.spinner("Processando o arquivo anexado pelas IAs auxiliares (isso pode levar alguns segundos)..."):
                    conteudo_extraido = process_file_attachment(uploaded_file, client)
                    if conteudo_extraido:
                        ideia_completa += f"\n\n{conteudo_extraido}"
            
            with st.spinner("Analisando o contexto e formulando perguntas investigativas..."):
                try:
                    chat_completion = client.chat.completions.create(
                        messages=[
                            {"role": "system", "content": get_system_prompt_triagem(tipo_ia)},
                            {"role": "user", "content": ideia_completa}
                        ],
                        model=MODEL_NAME,
                        temperature=0.7,
                    )
                    st.session_state.perguntas = chat_completion.choices[0].message.content
                    st.session_state.ideia_inicial = ideia_completa # Salvamos a ideia combinada com o anexo
                    st.session_state.tipo_ia = tipo_ia
                    st.session_state.step = 2
                    st.rerun()
                except Exception as e:
                    st.error(f"Erro na comunicação com a API: {e}")
        else:
            st.warning("Por favor, digite alguma ideia ou anexe um arquivo antes de continuar.")

# ETAPA 2: ENTREVISTA
elif st.session_state.step == 2:
    st.subheader("Etapa 2: Entrevista (Refinamento)")
    st.markdown(f"**🎯 Alvo:** {st.session_state.tipo_ia}\n\nPor favor, responda às perguntas abaixo para extrairmos o contexto perfeito para este modelo.")
    
    st.info(st.session_state.perguntas)
    
    respostas = st.text_area("Suas Respostas:", value=st.session_state.respostas, height=200, placeholder="1. O contexto é...\n2. O estilo será...")
    
    col1, col2 = st.columns([1, 1])
    with col1:
        if st.button("Voltar para Triagem"):
            st.session_state.respostas = respostas
            st.session_state.step = 1
            st.rerun()
    with col2:
        if st.button("Gerar Prompt Final", type="primary"):
            if respostas.strip():
                st.session_state.respostas = respostas
                with st.spinner("Sintetizando o seu Prompt Magistral..."):
                    contexto_completo = (
                        f"IDEIA INICIAL (INCLUINDO ANEXOS LIDOS):\n{st.session_state.ideia_inicial}\n\n"
                        f"PERGUNTAS DA ENTREVISTA:\n{st.session_state.perguntas}\n\n"
                        f"RESPOSTAS DO USUÁRIO:\n{st.session_state.respostas}"
                    )
                    try:
                        chat_completion = client.chat.completions.create(
                            messages=[
                                {"role": "system", "content": get_system_prompt_sintese(st.session_state.tipo_ia)},
                                {"role": "user", "content": contexto_completo}
                            ],
                            model=MODEL_NAME,
                            temperature=0.7,
                        )
                        st.session_state.prompt_final = chat_completion.choices[0].message.content
                        st.session_state.step = 3
                        st.rerun()
                    except Exception as e:
                        st.error(f"Erro na comunicação com a API: {e}")
            else:
                st.warning("Por favor, forneça as respostas antes de gerar o prompt.")

# ETAPA 3: SÍNTESE
elif st.session_state.step == 3:
    st.subheader("Etapa 3: Síntese (O Prompt Magistral)")
    st.markdown(f"Aqui está o seu prompt perfeitamente estruturado para **{st.session_state.tipo_ia}**.")
    
    st.code(st.session_state.prompt_final, language="markdown")
    
    st_copy_to_clipboard(
        text=st.session_state.prompt_final,
        before_copy_label="📋 Copiar para Área de Transferência",
        after_copy_label="✅ Copiado!"
    )
    
    st.write("") # Espaçamento
    col1, col2 = st.columns([1, 1])
    with col1:
        if st.button("Refazer Entrevista"):
            st.session_state.step = 2
            st.rerun()
    with col2:
        if st.button("Iniciar Novo Projeto", type="primary"):
            reset_app()
            st.rerun()
